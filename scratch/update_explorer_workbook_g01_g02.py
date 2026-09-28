from pathlib import Path

from openpyxl import load_workbook

path = Path(r"Scram Book/Explorer Tab/Rovvy_Explorer_Scrum_Book.xlsx")
wb = load_workbook(path)
ws = wb.active

h9 = (
    "2026-09-24: GET /api/v1/explore/places -> ExplorePlaceSpineService bounded SQL (LIMIT :lim; "
    "spatial ORDER BY distance/confidence/gers_id; city ORDER BY confidence/name/gers_id). "
    "migrations/007_places_city_category_explore_idx.sql. No FS/OSM route fallback; hub lat/lon/radius_m/limit + gers_id. "
    "pytest 42+ focused backend; Explore Vitest 103/17; tsc 0. Supabase migration apply + live data load not run here."
)
h10 = (
    "2026-09-24: migrations/006_place_ingest_runs.sql (+ place_ingest_run_cities coverage); "
    "04_load/05_resync begin ingest before mutations, fail closed on errors, idempotent succeeded scope; "
    "freshness joins city coverage (no NULL city_slug wildcard). 06_publish_places_r2 fail-closed + DuckDB parquet row-count verify. "
    "No production R2 upload in tests."
)

ws.cell(9, 4).value = "Complete in code"
ws.cell(9, 8).value = h9
ws.cell(10, 4).value = "Complete in code"
ws.cell(10, 8).value = h10
wb.save(path)
print("D9", ws.cell(9, 4).value)
print("H9", ws.cell(9, 8).value[:120], "...")
print("D10", ws.cell(10, 4).value)
print("H10", ws.cell(10, 8).value[:120], "...")
