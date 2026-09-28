from pathlib import Path

from openpyxl import load_workbook

path = Path(r"Scram Book/Explorer Tab/Rovvy_Explorer_Scrum_Book.xlsx")
wb = load_workbook(path)
ws = wb.active

h16 = (
    "2026-09-24: explore-listing-field-state.ts unified unknown copy: Price unknown, Photo unavailable, "
    "Hours unknown, Check provider, Sold out; no Unsplash stock for inventory (ExploreCardImage fallbackMode=unknown). "
    "Hub/category/drawer/event detail + picks table; View provider when price unknown + URL. "
    "Removed pseudoRating export from explore-events.ts (Explore scope). Global /events unchanged. "
    "ExploreCardImage useEffect resets loadFailed on imageUrl/placeId/fallbackMode change; "
    "explore-card-image-lifecycle.test.tsx. Explore Vitest: 101 passed (16 files); tsc 0; no new API."
)
h40 = (
    "2026-09-24: Complete in code for app/(dashboard)/explore only: no stars, review counts, dash ratings, "
    "or pseudoRating on hub/category/drawer/event detail/picks. explore-unknown-field-states.test.ts + "
    "explore-no-synthetic-ratings.test.ts + explore-card-image-lifecycle.test.tsx (photo error recovery). "
    "frontend/app/(dashboard)/events still uses local pseudoRating (out of scope)."
)

ws.cell(16, 4).value = "Complete in code"
ws.cell(16, 8).value = h16
ws.cell(40, 4).value = "Complete in code"
ws.cell(40, 8).value = h40
wb.save(path)

print("D16:", ws.cell(16, 4).value)
print("H16:", ws.cell(16, 8).value)
print("D40:", ws.cell(40, 4).value)
print("H40:", ws.cell(40, 8).value)
