"""F48 Explore drawer Directions → Live tab handoff (Complete in code, green)."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

path = Path(r"Scram Book/Explorer Tab/Rovvy_Explorer_Scrum_Book.xlsx")
wb = load_workbook(path)
ws = wb["Explorer Tasks"]
log = wb["Cursor Context Log"]

row = 74
f48_task = "[2026-10-05] Drawer map / directions"
evidence_add = (
    "2026-10-05 Explore Directions → /live?gers_id&lat&lng&name (Next Link, no Google Maps). "
    "live-explore-deeplink.ts build/parse; LivePageClient parseLiveDeepLink + selectDestination + router.replace(/live). "
    "UUID Overture gers_id accepted. Browser localhost:3000: Center for Native Futures → Live place panel + route preview "
    "(5.6 mi / 14 min); /live?lat=abc loads normal Live; Vitest explore+live 132+341 pass; tsc clean except ProfileTravelMap zIndexOffset."
)

ws.cell(row, 3).value = f48_task
ws.cell(row, 4).value = "Complete in code"
ws.cell(row, 7).value = (
    "Drawer Directions opens Rovvy Live with same place panel + route preview from user location; "
    "events use lat/lng/name only; hide when no coordinates."
)
old_ev = ws.cell(row, 8).value or ""
if evidence_add not in str(old_ev):
    ws.cell(row, 8).value = evidence_add + " " + str(old_ev).strip()

green = PatternFill(fill_type="solid", fgColor="DDF3E3")
ws.cell(row, 3).fill = green
ws.cell(row, 4).fill = green

plan = (
    "Cursor: regression — live-explore-deeplink Vitest; manual Explore drawer Directions → Live Back to Explore; "
    "no exploreDirectionsUrl / Google Maps."
)
ws.cell(row, 10).value = plan

log_row = log.max_row + 1
now = datetime(2026, 10, 5, 9, 45)
for c, val in enumerate(
    [
        now,
        "F48",
        "Capability",
        f48_task,
        "Complete in code",
        ws.cell(row, 5).value,
        ws.cell(row, 6).value,
        ws.cell(row, 7).value,
        ws.cell(row, 8).value,
        ws.cell(row, 9).value or "",
        ws.cell(row, 11).value or "Feature",
        "",
        plan,
        "Explorer",
    ],
    start=1,
):
    log.cell(log_row, c).value = val

wb.save(path)
print("F48 row", row, "updated; log row", log_row)
