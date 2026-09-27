"""Update F34 row on Explorer Tasks sheet after persistent save wiring."""
from pathlib import Path

from openpyxl import load_workbook

path = Path(r"Scram Book/Explorer Tab/Rovvy_Explorer_Scrum_Book.xlsx")
wb = load_workbook(path)
ws = wb.active

f34_evidence = (
    "2026-09-25 F34 (Partial in product): POST/GET /collection/items; saved_from "
    "explore:listing:place:{gers_id}|event:{id}; auth-gated reload; idempotent create_item; "
    "drawer save states + reload failure banner with Retry on /explore; signed-out → login. "
    "Vitest Explore 113/20; pytest collection 6; tsc 0. "
    "Manual QA still required: save→hard refresh→card+drawer Saved; account switch isolation; "
    "reload Retry when GET fails. No Explore unsave; coords often omitted on save body."
)

ws.cell(60, 4).value = "Partial in product"
ws.cell(60, 8).value = f34_evidence

wb.save(path)
print("F34 D60:", ws.cell(60, 4).value)
print("F34 H60:", ws.cell(60, 8).value[:100], "...")
