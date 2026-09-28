"""Add UI availability summary row + refresh as-of line on Explorer Tasks sheet."""
from pathlib import Path

from openpyxl import load_workbook

path = Path(r"Scram Book/Explorer Tab/Rovvy_Explorer_Scrum_Book.xlsx")
wb = load_workbook(path)
ws = wb.active

ws.cell(3, 1).value = (
    "As of 24 Sep 2026 — 76 tasks; 33 Complete in code, 23 Partial, 15 Pending, 5 Not started "
    "(status column D). “Complete in code” is not production or browser QA."
)

ui_evidence = (
    "2026-09-24 audit: G08 unknown-state copy — working in code (Vitest explore-unknown-field-states, "
    "explore-availability-copy). F15 open-now/capacity/walk-in — not implemented; Overture place cards "
    "show Hours unknown / Check provider unless spine hours exist; Ticketmaster Sold out when provider says so. "
    "UI functions: normalizeListingAvailability, placeCardAvailability, placeHoursLabel, hubListingBadge, "
    "drawerHoursSourceLabel, categoryCardScheduleLine; exploreEventPriceState, exploreSlotPriceLabel, "
    "exploreListingPhotoState, exploreInventoryPhotoDisplay, exploreDrawerHoursState, exploreAvailabilityFieldState, "
    "exploreDrawerProviderActionLabel, exploreSlotMatchesNumericPriceChip; wired via eventToSlot/placeToSlot → "
    "ExploreDetailDrawer + category cards + ExploreCardImage fallbackMode=unknown. Explore Vitest 103/17; tsc 0."
)

ws.cell(4, 1).value = "UI availability"
ws.cell(4, 2).value = "G08 working · F15 partial"
ws.cell(4, 3).value = "Unknown price/hours/photo/availability UI"
ws.cell(4, 4).value = "Partial in product"
ws.cell(4, 5).value = "Critical"
ws.cell(4, 6).value = "Medium"
ws.cell(4, 7).value = (
    "Do not show open-now, capacity, or bookable without verified provider/spine fields; keep G08 honest copy."
)
ws.cell(4, 8).value = ui_evidence
ws.cell(4, 9).value = "G08, F15, F13, F11"

wb.save(path)
print("Row 3:", ws.cell(3, 1).value[:80], "...")
print("Row 4 D:", ws.cell(4, 4).value)
